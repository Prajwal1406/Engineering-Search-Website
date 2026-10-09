#!/usr/bin/env python3
"""Engineering Search V2: row-level Excel/CSV index with SQLite FTS5.

- Source documents are opened read-only.
- One database record is stored per meaningful spreadsheet row, not per cell.
- Incremental scans skip unchanged files.
- FTS5 trigram search supports model numbers and Japanese substring searches.
"""

# Select-Object Name, @{Name="SizeMB";Expression={:Round($_.Length / 1MB, 1)}}
from __future__ import annotations

import argparse
import csv
import os
import re
import sqlite3
import subprocess
import threading
import time
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Missing openpyxl. Run: py -m pip install openpyxl")
    raise

DEFAULT_ROOT = os.getenv("ENGINEERING_SEARCH_ROOT", str(Path(__file__).resolve().parent / "sample_documents"))
DEFAULT_DB = os.getenv("ENGINEERING_SEARCH_DB", str(Path(__file__).resolve().parent / "databases" / "engineering_search.db"))
SUPPORTED = {".xlsx", ".xlsm", ".csv", ".tsv"}
EXCLUDED_DIRS = {"$RECYCLE.BIN", "System Volume Information", ".git", "__pycache__"}
OLD_WORDS = {"旧", "旧版", "old", "obsolete", "backup", "bak"}
BOM_WORDS = {"部品表", "部品リスト", "パーツリスト", "bom", "parts list", "購入品", "機器一覧", "lst"}
MAX_ROW_TEXT = 4000
INSERT_BATCH = 2000
COMMIT_EVERY_FILES = 25


def open_db(path: Path, readonly: bool = False) -> sqlite3.Connection:
    if readonly:
        uri = path.resolve().as_uri() + "?mode=ro"
        con = sqlite3.connect(uri, uri=True, timeout=30)
        con.execute("PRAGMA query_only=ON")
        return con
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(path, timeout=60)
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA synchronous=NORMAL")
    con.execute("PRAGMA foreign_keys=ON")
    return con


def initialize(con: sqlite3.Connection) -> None:
    con.executescript("""
    CREATE TABLE IF NOT EXISTS files(
        id INTEGER PRIMARY KEY,
        full_path TEXT NOT NULL UNIQUE,
        file_name TEXT NOT NULL,
        extension TEXT NOT NULL,
        relative_path TEXT NOT NULL,
        parent_folder TEXT NOT NULL,
        modified_ns INTEGER NOT NULL,
        size_bytes INTEGER NOT NULL,
        is_old INTEGER NOT NULL DEFAULT 0,
        is_bom INTEGER NOT NULL DEFAULT 0,
        status TEXT NOT NULL DEFAULT 'active',
        row_count INTEGER NOT NULL DEFAULT 0,
        indexed_at TEXT NOT NULL,
        error_text TEXT
    );
    CREATE TABLE IF NOT EXISTS rows_data(
        id INTEGER PRIMARY KEY,
        file_id INTEGER NOT NULL,
        sheet_name TEXT NOT NULL,
        row_number INTEGER NOT NULL,
        row_text TEXT NOT NULL,
        FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
    );
    CREATE INDEX IF NOT EXISTS idx_files_path ON files(full_path);
    CREATE INDEX IF NOT EXISTS idx_rows_file ON rows_data(file_id);
    CREATE TABLE IF NOT EXISTS scan_log(
        id INTEGER PRIMARY KEY,
        started_at TEXT,
        finished_at TEXT,
        root_path TEXT,
        checked INTEGER DEFAULT 0,
        new_files INTEGER DEFAULT 0,
        updated_files INTEGER DEFAULT 0,
        unchanged_files INTEGER DEFAULT 0,
        error_files INTEGER DEFAULT 0,
        indexed_rows INTEGER DEFAULT 0
    );
    """)
    # Trigram is appropriate for model and Japanese substring search.
    try:
        con.execute("""CREATE VIRTUAL TABLE IF NOT EXISTS rows_fts USING fts5(
            row_text,
            content='rows_data',
            content_rowid='id',
            tokenize='trigram case_sensitive 0'
        )""")
    except sqlite3.OperationalError as exc:
        if "tokenizer" in str(exc).lower():
            raise RuntimeError("This Python SQLite build does not support the FTS5 trigram tokenizer.") from exc
        raise
    con.executescript("""
    CREATE TRIGGER IF NOT EXISTS rows_ai AFTER INSERT ON rows_data BEGIN
      INSERT INTO rows_fts(rowid,row_text) VALUES(new.id,new.row_text);
    END;
    CREATE TRIGGER IF NOT EXISTS rows_ad AFTER DELETE ON rows_data BEGIN
      INSERT INTO rows_fts(rows_fts,rowid,row_text) VALUES('delete',old.id,old.row_text);
    END;
    CREATE TRIGGER IF NOT EXISTS rows_au AFTER UPDATE ON rows_data BEGIN
      INSERT INTO rows_fts(rows_fts,rowid,row_text) VALUES('delete',old.id,old.row_text);
      INSERT INTO rows_fts(rowid,row_text) VALUES(new.id,new.row_text);
    END;
    """)
    con.commit()


def path_flags(path: Path) -> tuple[int, int]:
    parts = [part.casefold() for part in path.parts]
    joined = " ".join(parts)
    return int(any(word in parts for word in OLD_WORDS)), int(any(word in joined for word in BOM_WORDS))


def supported_files(root: Path):
    def onerror(err):
        print(f"[Skipped folder] {err}", flush=True)
    for current, dirs, names in os.walk(root, onerror=onerror):
        dirs[:] = [d for d in dirs if d not in EXCLUDED_DIRS]
        for name in names:
            if name.startswith(("~$", ".~lock.")):
                continue
            path = Path(current) / name
            if path.suffix.lower() in SUPPORTED:
                yield path


def normalize_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def meaningful_row(values: list[str]) -> bool:
    if not values:
        return False
    text = " ".join(values).strip()
    if not text:
        return False
    # Exclude rows made only of punctuation/box marks.
    if re.fullmatch(r"[\s\-―ー_=|/\\.,:;・○●◎△▲▽▼×※*]+", text):
        return False
    # Exclude a single plain number/date-like serial, while preserving engineering codes.
    if len(values) == 1 and re.fullmatch(r"[\d,./:\- ]+", values[0]):
        return False
    return True


def spreadsheet_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for ws in wb.worksheets:
            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = [normalize_value(v) for v in row]
                values = [v for v in values if v]
                if meaningful_row(values):
                    text = " | ".join(values)
                    yield ws.title, row_no, text[:MAX_ROW_TEXT]
    finally:
        wb.close()


def delimited_rows(path: Path):
    data = path.read_bytes()
    decoded = None
    for enc in ("utf-8-sig", "cp932", "utf-16"):
        try:
            decoded = data.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        decoded = data.decode("utf-8", errors="replace")
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    for row_no, row in enumerate(csv.reader(decoded.splitlines(), delimiter=delimiter), start=1):
        values = [normalize_value(v) for v in row]
        values = [v for v in values if v]
        if meaningful_row(values):
            yield "CSV", row_no, " | ".join(values)[:MAX_ROW_TEXT]


def index_file(con: sqlite3.Connection, root: Path, path: Path) -> tuple[str, int]:
    st = path.stat()
    existing = con.execute("SELECT id,modified_ns,size_bytes FROM files WHERE full_path=?", (str(path),)).fetchone()
    if existing and existing[1] == st.st_mtime_ns and existing[2] == st.st_size:
        con.execute("UPDATE files SET status='active',error_text=NULL WHERE id=?", (existing[0],))
        return "unchanged", 0

    is_old, is_bom = path_flags(path)
    now = datetime.now().isoformat(timespec="seconds")
    if existing:
        file_id = existing[0]
        con.execute("DELETE FROM rows_data WHERE file_id=?", (file_id,))
        con.execute("""UPDATE files SET file_name=?,extension=?,relative_path=?,parent_folder=?,
            modified_ns=?,size_bytes=?,is_old=?,is_bom=?,status='active',indexed_at=?,error_text=NULL,row_count=0
            WHERE id=?""", (path.name,path.suffix.lower(),str(path.relative_to(root)),path.parent.name,
            st.st_mtime_ns,st.st_size,is_old,is_bom,now,file_id))
        state = "updated"
    else:
        cur = con.execute("""INSERT INTO files(full_path,file_name,extension,relative_path,parent_folder,
            modified_ns,size_bytes,is_old,is_bom,status,indexed_at)
            VALUES(?,?,?,?,?,?,?,?,?,'active',?)""", (str(path),path.name,path.suffix.lower(),
            str(path.relative_to(root)),path.parent.name,st.st_mtime_ns,st.st_size,is_old,is_bom,now))
        file_id = cur.lastrowid
        state = "new"

    iterator = delimited_rows(path) if path.suffix.lower() in {".csv", ".tsv"} else spreadsheet_rows(path)
    batch, count = [], 0
    for sheet, row_no, text in iterator:
        batch.append((file_id, sheet, row_no, text))
        count += 1
        if len(batch) >= INSERT_BATCH:
            con.executemany("INSERT INTO rows_data(file_id,sheet_name,row_number,row_text) VALUES(?,?,?,?)", batch)
            batch.clear()
    if batch:
        con.executemany("INSERT INTO rows_data(file_id,sheet_name,row_number,row_text) VALUES(?,?,?,?)", batch)
    con.execute("UPDATE files SET row_count=? WHERE id=?", (count, file_id))
    return state, count


def scan(root: Path, db: Path) -> None:
    if not root.is_dir():
        raise FileNotFoundError(f"Folder not found: {root}")
    con = open_db(db)
    initialize(con)
    started = datetime.now().isoformat(timespec="seconds")
    log_id = con.execute("INSERT INTO scan_log(started_at,root_path) VALUES(?,?)", (started,str(root))).lastrowid
    con.execute("UPDATE files SET status='unseen' WHERE full_path LIKE ?", (str(root) + "%",))
    stats = {"checked":0,"new":0,"updated":0,"unchanged":0,"errors":0,"rows":0}
    print(f"V2 row-level scan\nRoot: {root}\nDatabase: {db}\n", flush=True)

    for path in supported_files(root):
        stats["checked"] += 1
        try:
            state, row_count = index_file(con, root, path)
            stats[state] += 1
            stats["rows"] += row_count
            print(f"[{stats['checked']:,}] {state.upper():9} rows={row_count:,}  {path}", flush=True)
        except Exception as exc:
            stats["errors"] += 1
            print(f"[{stats['checked']:,}] ERROR {path}: {exc}", flush=True)
            try:
                st = path.stat(); old,bom = path_flags(path); now = datetime.now().isoformat(timespec="seconds")
                con.execute("""INSERT INTO files(full_path,file_name,extension,relative_path,parent_folder,modified_ns,size_bytes,is_old,is_bom,status,indexed_at,error_text)
                VALUES(?,?,?,?,?,?,?,?,?,'error',?,?) ON CONFLICT(full_path) DO UPDATE SET status='error',indexed_at=excluded.indexed_at,error_text=excluded.error_text""",
                (str(path),path.name,path.suffix.lower(),str(path.relative_to(root)),path.parent.name,st.st_mtime_ns,st.st_size,old,bom,now,str(exc)))
            except Exception:
                pass
        if stats["checked"] % COMMIT_EVERY_FILES == 0:
            con.commit()
            print(f"  Saved progress: files={stats['checked']:,}, indexed rows={stats['rows']:,}", flush=True)

    missing = con.execute("SELECT COUNT(*) FROM files WHERE status='unseen' AND full_path LIKE ?", (str(root)+"%",)).fetchone()[0]
    con.execute("UPDATE files SET status='missing' WHERE status='unseen' AND full_path LIKE ?", (str(root)+"%",))
    finished = datetime.now().isoformat(timespec="seconds")
    con.execute("""UPDATE scan_log SET finished_at=?,checked=?,new_files=?,updated_files=?,unchanged_files=?,error_files=?,indexed_rows=? WHERE id=?""",
                (finished,stats['checked'],stats['new'],stats['updated'],stats['unchanged'],stats['errors'],stats['rows'],log_id))
    con.commit()
    con.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    con.close()
    print("\nV2 scan completed", flush=True)
    print(f"Checked: {stats['checked']:,} | New: {stats['new']:,} | Updated: {stats['updated']:,} | Unchanged: {stats['unchanged']:,} | Errors: {stats['errors']:,} | Missing: {missing:,}", flush=True)
    print(f"New/updated meaningful rows: {stats['rows']:,}", flush=True)


def fts_query(raw: str) -> str:
    terms = [t.strip() for t in raw.split() if t.strip()]
    # Quote each term. Trigram supports substring matching for terms of 3+ characters.
    return " AND ".join('"' + t.replace('"', '""') + '"' for t in terms)


def gui(db: Path) -> None:
    import tkinter as tk
    from tkinter import ttk, messagebox

    class App(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title("Engineering Search V2")
            self.geometry("1400x760")
            self.db = db
            self.searching = False

            bar = ttk.Frame(self,padding=10); bar.pack(fill="x")
            ttk.Label(bar,text="Part / keyword:").pack(side="left")
            self.entry = ttk.Entry(bar,width=48); self.entry.pack(side="left",padx=8)
            self.entry.bind("<Return>",lambda e:self.start_search())
            self.bom = tk.BooleanVar(value=True)
            ttk.Checkbutton(bar,text="Prioritize BOM / 部品表",variable=self.bom).pack(side="left",padx=8)
            self.search_btn = ttk.Button(bar,text="Search",command=self.start_search); self.search_btn.pack(side="left")
            ttk.Button(bar,text="Clear",command=self.clear).pack(side="left",padx=6)
            self.status = ttk.Label(self,text=f"Database: {db}",padding=(10,0)); self.status.pack(fill="x")
            self.progress = ttk.Progressbar(self,mode="indeterminate")

            cols=("file","sheet","row","evidence","bom","old","modified","path")
            self.tree=ttk.Treeview(self,columns=cols,show="headings")
            labels={"file":"File","sheet":"Sheet","row":"Row","evidence":"Matching row content","bom":"BOM?","old":"Old?","modified":"Modified","path":"Full path"}
            widths={"file":230,"sheet":120,"row":65,"evidence":430,"bom":55,"old":55,"modified":140,"path":500}
            for col in cols:
                self.tree.heading(col,text=labels[col]); self.tree.column(col,width=widths[col],anchor="w")
            self.tree.pack(fill="both",expand=True,padx=10,pady=8)
            self.tree.bind("<Double-1>",lambda e:self.open_file())
            buttons=ttk.Frame(self,padding=10); buttons.pack(fill="x")
            ttk.Button(buttons,text="Open File",command=self.open_file).pack(side="left")
            ttk.Button(buttons,text="Open Folder",command=self.open_folder).pack(side="left",padx=6)
            ttk.Button(buttons,text="Copy Path",command=self.copy_path).pack(side="left")

        def selected_path(self):
            selected=self.tree.selection()
            return self.tree.item(selected[0],"values")[-1] if selected else None
        def clear(self):
            self.entry.delete(0,"end")
            for i in self.tree.get_children(): self.tree.delete(i)
            self.status.config(text="Enter a model or keyword.")
        def start_search(self):
            if self.searching: return
            raw=self.entry.get().strip()
            if not raw: return
            if any(len(t)<3 for t in raw.split()):
                messagebox.showinfo("Search term too short","V2 fast search requires each term to contain at least 3 characters. Examples: QY80, PS21, PLC, シーケンサ")
                return
            self.searching=True; self.search_btn.config(state="disabled"); self.entry.config(state="disabled")
            self.status.config(text=f'Searching for "{raw}"...')
            self.progress.pack(fill="x",padx=10,pady=(4,0)); self.progress.start(12)
            threading.Thread(target=self.run_search,args=(raw,self.bom.get()),daemon=True).start()
        def run_search(self,raw,prioritize):
            started=time.perf_counter()
            try:
                con=open_db(self.db,readonly=True)
                sort="f.is_bom DESC," if prioritize else ""
                sql=f"""SELECT f.file_name,r.sheet_name,r.row_number,r.row_text,f.is_bom,f.is_old,
                                 datetime(f.modified_ns/1000000000,'unixepoch','localtime'),f.full_path
                          FROM rows_fts x JOIN rows_data r ON r.id=x.rowid JOIN files f ON f.id=r.file_id
                          WHERE rows_fts MATCH ? AND f.status='active'
                          ORDER BY {sort} f.is_old ASC,f.modified_ns DESC LIMIT 5000"""
                rows=con.execute(sql,(fts_query(raw),)).fetchall(); con.close()
                self.after(0,self.finish,rows,time.perf_counter()-started,None)
            except Exception as exc:
                self.after(0,self.finish,[],time.perf_counter()-started,str(exc))
        def finish(self,rows,elapsed,error):
            self.progress.stop(); self.progress.pack_forget(); self.search_btn.config(state="normal"); self.entry.config(state="normal"); self.searching=False
            for i in self.tree.get_children(): self.tree.delete(i)
            if error:
                self.status.config(text="Search failed."); messagebox.showerror("Search error",error); return
            for row in rows:
                out=list(row); out[4]="Yes" if row[4] else "No"; out[5]="Yes" if row[5] else "No"
                self.tree.insert("","end",values=out)
            self.status.config(text=f"{len(rows):,} matching row(s) found in {elapsed:.2f} seconds.")
            self.entry.focus_set()
        def open_file(self):
            path=self.selected_path()
            if path and Path(path).exists(): os.startfile(path)
            elif path: messagebox.showwarning("Unavailable",f"File is unavailable:\n{path}")
        def open_folder(self):
            path=self.selected_path()
            if path: subprocess.Popen(["explorer","/select,",path])
        def copy_path(self):
            path=self.selected_path()
            if path: self.clipboard_clear(); self.clipboard_append(path)

    if not db.exists():
        raise FileNotFoundError(f"Database not found: {db}")
    App().mainloop()


def main() -> int:
    parser=argparse.ArgumentParser(description="Engineering Search V2 row-level SQLite FTS indexer")
    parser.add_argument("--root",default=DEFAULT_ROOT)
    parser.add_argument("--db",default=DEFAULT_DB)
    parser.add_argument("--scan",action="store_true")
    parser.add_argument("--gui",action="store_true")
    args=parser.parse_args(); root=Path(args.root); db=Path(args.db)
    if not args.scan and not args.gui:
        parser.print_help(); return 0
    try:
        if args.scan: scan(root,db)
        if args.gui: gui(db)
    except KeyboardInterrupt:
        print("\nStopped by user. Committed batches remain available for the next incremental run.")
        return 130
    except Exception as exc:
        print(f"ERROR: {exc}")
        return 1
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
